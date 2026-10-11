# -*- coding: utf-8 -*-
"""把官方导出的期刊分区源表，合并成本项目用的本地映射表。

用法
    python scripts/pipeline/build_journal_ranks.py <源表.xlsx> [-o 输出路径]
    （不给 -o 时输出到 webapp/config/journal_ranks.csv）

为什么需要自己合并
    分区表是**授权数据**（科睿唯安 JCR / 中科院分区 / 新锐分区），本项目**不分发**它，
    也不分发合并结果（webapp/config/journal_ranks.csv 已在 .gitignore 里）。
    你从正规渠道拿到官方 xlsx 后，在本机跑这个脚本生成自己的表即可。

源表结构（4 个工作表；**表头不在第 1 行**，脚本自动定位含关键词的那一行）
    2026新锐期刊分区表        刊名 | ISSN | EISSN | 新锐分区 | TOP        ← 「英文名/中文名」是**学科**，不是刊名
    2025年版中科院期刊分区表   期刊名称 | 2025分区 | Top | Open Access    ← 无 ISSN
    2023 vs 2025分区对比      期刊名称 | 2025分区 | 2023分区 | Top | Open Access
    2026年度JCR 期刊分区       期刊名称 | ISSN | eISSN | 学科类别 | 影响因子JIF | JIF分区 | JIF百分位 | …

产出列（与 core/journal_ranks.py 的 FIELDS 对齐）
    刊名 ISSN eISSN 新锐分区 新锐TOP 新锐学科 中科院2025 中科院2023
    中科院TOP 中科院OA JCR分区 JIF JIF百分位 JCR学科 来源 别名

匹配与消歧
    * 统一 key = 去掉所有非字母数字并转大写（源表刊名有三套写法：新锐去空格、JCR 全大写、
      中科院正常空格）。能拿到 ISSN 时优先按 ISSN 精确匹配，刊名只作兜底。
    * **同一张表内**，同一个 loose key 对应多个不同原始刊名时（实测 18 组，例如
      Science & Education 与 SCIENCE EDUCATION 是两本不同的刊）**不许按刊名合并**，
      只按 ISSN —— 否则张冠李戴。
"""
from __future__ import annotations

import argparse
import csv
import os
import re
import sys

import openpyxl

# ----------------------------------------------------------------- 基础工具
OUT_FIELDS = ["刊名", "ISSN", "eISSN", "新锐分区", "新锐TOP", "新锐学科",
              "中科院2025", "中科院2023", "中科院TOP", "中科院OA",
              "JCR分区", "JIF", "JIF百分位", "JCR学科", "来源", "别名"]

# 源表 → 输出用的字段名（值就是源表里的列标题）
HEADER_KEYS = ("刊名", "期刊名称")          # 用来定位表头行


def norm_name(s):
    """刊名 key：去掉所有非字母数字 + 转大写（抹平三套源表的写法差异）。"""
    return re.sub(r"[^A-Z0-9]", "", str(s or "").upper())


def norm_issn(s):
    v = re.sub(r"[^0-9Xx]", "", str(s or "")).upper()
    return v if len(v) == 8 else ""


def _s(v):
    """单元格取值 → 去空白字符串（None/数字都安全）。"""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def find_header(ws, max_scan=8):
    """定位表头行（列标题里含『刊名』或『期刊名称』），返回 (0-based 行号, [列标题])。"""
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=max_scan, values_only=True)):
        cells = [_s(c) for c in row]
        if any(k in c for c in cells for k in HEADER_KEYS):
            return i, cells
    raise RuntimeError(f"工作表「{ws.title}」里找不到表头行（前 {max_scan} 行内没有『刊名/期刊名称』）")


def read_sheet(ws):
    """把一张工作表读成 [dict]，键是列标题。表头行以下的数据行都读。"""
    hrow, headers = find_header(ws)
    out = []
    for row in ws.iter_rows(min_row=hrow + 2, values_only=True):   # +1 转 1-based，+1 跳过表头
        vals = [_s(c) for c in row]
        if not any(vals):
            continue
        rec = {h: vals[i] for i, h in enumerate(headers) if i < len(vals) and h}
        if any(rec.get(k) for k in HEADER_KEYS):    # 有些表尾部带小标题行，丢掉无刊名的
            out.append(rec)
    return out


def pick(rec, *names):
    """从 rec 里按候选列名取第一个非空值（源表列名偶有细微差异时兜底）。"""
    for n in names:
        v = rec.get(n)
        if v:
            return v
    return ""


def ambiguous_keys(records, name_key):
    """**单张表内部**的歧义 key：loose key 相同、但原始刊名不同。

    例：Science & Education（Springer）与 SCIENCE EDUCATION（Wiley）去掉非字母数字后
    同 key，却是两本刊——这类 key 一律不按刊名合并，只按 ISSN。

    ⚠ 必须**逐表**判定，不能把几张表的名字混在一起算：新锐用正常大小写、JCR 用全大写，
    同一本刊在两表里天然长得不一样，混算会把几乎所有刊都误判成歧义（踩过）。
    """
    names_for = {}
    for r in records:
        nm = _s(r.get(name_key))
        k = norm_name(nm)
        if k:
            names_for.setdefault(k, set()).add(nm)
    return {k for k, s in names_for.items() if len(s) > 1}


# ----------------------------------------------------------------- 合并
class Reg:
    """一张刊的表：既按 ISSN 索引，也按归一刊名索引。

    ambiguous 记录「同一张表内、loose key 相同但原始刊名不同」的 key ——
    这些 key 一律不许按刊名合并。
    """

    def __init__(self):
        self.rows = []            # [{key, issn, name, ...}]
        self.by_issn = {}
        self.by_name = {}
        self._names_for = {}

    def add(self, name, issn, payload):
        k = norm_name(name)
        i = norm_issn(issn)
        rec = {"key": k, "issn": i, "name": _s(name), **payload}
        self.rows.append(rec)
        if i:
            self.by_issn.setdefault(i, rec)
        if k:
            self._names_for.setdefault(k, set()).add(_s(name))
            self.by_name.setdefault(k, rec)
        return rec

    @property
    def ambiguous(self):
        return {k for k, s in self._names_for.items() if len(s) > 1}

    def find(self, name, issn):
        i = norm_issn(issn)
        if i and i in self.by_issn:
            return self.by_issn[i]
        k = norm_name(name)
        if k and k not in self.ambiguous and k in self.by_name:
            return self.by_name[k]
        return None


def build(src_path):
    wb = openpyxl.load_workbook(src_path, read_only=True, data_only=True)
    sheets = {}
    for ws in wb.worksheets:
        try:
            sheets[ws.title] = read_sheet(ws)
        except RuntimeError:
            sheets[ws.title] = []          # 缺表头的表直接跳过，不阻断其他表
    wb.close()

    def sheet_like(*frags):
        for t, rows in sheets.items():
            if all(f in t for f in frags):
                return rows
        return []

    xr = sheet_like("新锐")                       # 新锐
    cas = sheet_like("中科院")                    # 中科院（2025 版）
    cmp_ = sheet_like("对比")                     # 2023 vs 2025 对比
    jcr = sheet_like("JCR")                       # JCR
    print("读入：新锐 %d / 中科院 %d / 对比 %d / JCR %d" % (len(xr), len(cas), len(cmp_), len(jcr)))

    # 每张表内部的歧义 key（必须逐表判定，见 ambiguous_keys 注释）
    AMB = {
        "xr":  ambiguous_keys(xr, "刊名"),
        "jcr": ambiguous_keys(jcr, "期刊名称"),
        "cas": ambiguous_keys(cas, "期刊名称"),
        "cmp": ambiguous_keys(cmp_, "期刊名称"),
    }

    # 输出表：以「JCR + 新锐」为骨架（这两张带 ISSN），再用中科院 / 对比按刊名补分区。
    # JCR 放前面：它的刊名是「全大写 + 正常空格」，作为显示名最可读（旧表也是这个写法）。
    merged = []            # 输出行
    by_issn = {}           # 归一 ISSN → 输出行
    by_key = {}            # 归一刊名 → 输出行

    def blank():
        return {k: "" for k in OUT_FIELDS} | {"_flags": set()}

    def slot(name, issn, amb):
        """按 ISSN 优先找到（或新建）输出行；无 ISSN / 匹配不上时按刊名兜底（歧义 key 除外）。"""
        i = norm_issn(issn)
        k = norm_name(name)
        row = by_issn.get(i) if i else None
        if row is None and k and k not in amb:
            row = by_key.get(k)
        if row is None:
            row = blank()
            row["刊名"] = _s(name)
            merged.append(row)
        if i:
            by_issn.setdefault(i, row)
            if not row["ISSN"]:
                row["ISSN"] = i
        if k and k not in amb and k not in by_key:
            by_key[k] = row
        return row

    # ① JCR（带 ISSN，定骨架）
    for r in jcr:
        row = slot(pick(r, "期刊名称", "刊名"), pick(r, "ISSN"), AMB["jcr"])
        row["JCR分区"] = pick(r, "JIF分区")
        row["JIF"] = pick(r, "影响因子JIF")
        row["JIF百分位"] = pick(r, "JIF百分位")
        row["JCR学科"] = pick(r, "学科类别")
        if not row["eISSN"]:
            row["eISSN"] = norm_issn(pick(r, "eISSN", "EISSN"))
        row["_flags"].add("jcr")

    # ② 新锐（带 ISSN）
    for r in xr:
        row = slot(pick(r, "刊名"), pick(r, "ISSN"), AMB["xr"])
        row["新锐分区"] = pick(r, "新锐分区")
        row["新锐TOP"] = pick(r, "TOP")
        row["新锐学科"] = pick(r, "中文名")          # 中文名 = 学科（不是刊名）
        if not row["eISSN"]:
            row["eISSN"] = norm_issn(pick(r, "EISSN"))
        row["_flags"].add("xr")

    # ③ 中科院（无 ISSN，只能按刊名匹配）
    for r in cas:
        row = slot(pick(r, "期刊名称", "刊名"), "", AMB["cas"])
        if not row["中科院2025"]:
            row["中科院2025"] = pick(r, "2025分区")
        if not row["中科院TOP"]:
            row["中科院TOP"] = pick(r, "Top")
        if not row["中科院OA"]:
            row["中科院OA"] = pick(r, "Open Access")
        row["_flags"].add("cas")

    # ④ 2023 vs 2025 对比（补 2023 分区）
    for r in cmp_:
        row = slot(pick(r, "期刊名称", "刊名"), "", AMB["cmp"])
        if not row["中科院2023"]:
            row["中科院2023"] = pick(r, "2023分区")
        if not row["中科院2025"]:
            row["中科院2025"] = pick(r, "2025分区")
        if not row["中科院TOP"]:
            row["中科院TOP"] = pick(r, "Top")
        if not row["中科院OA"]:
            row["中科院OA"] = pick(r, "Open Access")
        row["_flags"].add("cmp")

    # 收尾：来源字符串 + 别名
    for row in merged:
        row["来源"] = "+".join(sorted(row.pop("_flags")))
        row["别名"] = norm_name(row["刊名"])

    # 排序：中科院2025 → 新锐分区 → 刊名（与原表一致）
    def tier(v):
        m = re.search(r"\d", str(v or ""))
        return int(m.group()) if m else 9

    merged.sort(key=lambda r: (tier(r.get("中科院2025")), tier(r.get("新锐分区")), r.get("刊名", "")))
    return merged


def main(argv=None):
    ap = argparse.ArgumentParser(description="合并官方分区源表 → journal_ranks.csv")
    ap.add_argument("src", help="源表 .xlsx 路径（四张表那个文件）")
    ap.add_argument("-o", "--out", default=os.path.join(
        os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
        "webapp", "config", "journal_ranks.csv"), help="输出 CSV 路径")
    a = ap.parse_args(argv)

    if not os.path.isfile(a.src):
        print("源表不存在：%s" % a.src)
        return 1
    rows = build(a.src)
    os.makedirs(os.path.dirname(a.out), exist_ok=True)
    with open(a.out, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=OUT_FIELDS, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)
    hit = sum(1 for r in rows if r["中科院2025"] or r["新锐分区"] or r["JCR分区"])
    with_issn = sum(1 for r in rows if r["ISSN"])
    print("已写出：%s" % a.out)
    print("共 %d 条（带 ISSN %d 条，含至少一个分区 %d 条）" % (len(rows), with_issn, hit))
    print("提示：这是授权数据的本地映射表，不要提交或分发。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
